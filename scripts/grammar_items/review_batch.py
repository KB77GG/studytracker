"""Seeded, capacity-constrained sampling and a deterministic blind workbook."""

import random
import re
import unicodedata
from collections import Counter, defaultdict, deque
from datetime import datetime
from io import BytesIO
from math import ceil
from pathlib import Path
from zipfile import ZIP_STORED, ZipFile, ZipInfo

from .templates import TEMPLATES

HEADERS = (
    "序号",
    "题目编号",
    "知识点（ID + 中文名）",
    "题干",
    "句子",
    "选项",
    "机器答案",
    "老师判定",
    "正确答案应为（选填）",
    "备注",
)
VERDICTS = ("正确", "答案错", "题目有歧义", "句子不适合出题")
SHEET_NAMES = dict(
    zip(
        TEMPLATES,
        ("主句谓语", "从句边界", "从句类型", "状语关系", "完整主语", "动词身份"),
        strict=True,
    )
)
WIDTHS = (6, 23, 25, 34, 70, 32, 70, 24, 42, 32)


def largest_remainder(counts, total):
    denominator = sum(counts.values())
    if not denominator:
        return {}
    quotas = {key: count * total // denominator for key, count in sorted(counts.items())}
    order = sorted(counts, key=lambda key: (-(counts[key] * total % denominator), key))
    for key in order[: total - sum(quotas.values())]:
        quotas[key] += 1
    return quotas


def add_edge(graph, left, right, capacity):
    graph[left][right] = capacity
    graph[right].setdefault(left, 0)


def fill_flow(graph, source, sink):
    flow = 0
    while True:
        parents, queue = {source: None}, deque([source])
        while queue and sink not in parents:
            node = queue.popleft()
            for neighbor, capacity in graph[node].items():
                if capacity > 0 and neighbor not in parents:
                    parents[neighbor] = node
                    queue.append(neighbor)
        if sink not in parents:
            return flow
        node = sink
        while parents[node] is not None:
            previous = parents[node]
            graph[previous][node] -= 1
            graph[node][previous] += 1
            node = previous
        flow += 1


def constrained_sample(pool, quotas, rng, used_sentences, used_passages, stratified=True):
    groups = defaultdict(list)
    for item in sorted(pool, key=lambda x: x["item_id"]):
        source = item["source"]
        key = (source["passage_id"], source["sentence_id"])
        if key not in used_sentences and used_passages[key[0]] < 2:
            groups[key].append(item)
    ordered = list(sorted(groups))
    rng.shuffle(ordered)
    graph, choices = defaultdict(dict), {}
    start, end = ("start",), ("end",)
    for key, quota in sorted(quotas.items()):
        add_edge(graph, start, ("stratum", *key), quota)
    for key in ordered:
        item = rng.choice(groups[key])
        stratum = (item["source"]["source_kind"], item["difficulty"]) if stratified else ("all",)
        sentence, passage = ("sentence", *key), ("passage", key[0])
        add_edge(graph, ("stratum", *stratum), sentence, 1)
        add_edge(graph, sentence, passage, 1)
        choices[sentence] = item
    for pid in sorted({key[0] for key in groups}):
        add_edge(graph, ("passage", pid), end, 2 - used_passages[pid])
    wanted = sum(quotas.values())
    if fill_flow(graph, start, end) != wanted:
        raise ValueError(f"Cannot meet {wanted} review slots with strata/sentence/passage limits")
    selected = [item for node, item in choices.items() if graph[node][("passage", node[1])] == 0]
    for item in selected:
        source = item["source"]
        used_sentences.add((source["passage_id"], source["sentence_id"]))
        used_passages[source["passage_id"]] += 1
    return selected


def sample_batch(candidates, seed, *, auto_count=50, review_count=10):
    rng, batch, stats = random.Random(seed), [], {}
    for template in TEMPLATES:
        pool = [item for item in candidates if item["template"] == template]
        auto = [item for item in pool if item["status"] == "auto_pass"]
        review = [item for item in pool if item["status"] == "needs_review"]
        count = min(auto_count, len(auto))
        sizes = Counter((item["source"]["source_kind"], item["difficulty"]) for item in auto)
        quotas = largest_remainder(sizes, count)
        sentences, passages = set(), Counter()
        chosen = constrained_sample(auto, quotas, rng, sentences, passages)
        chosen += constrained_sample(
            review, {("all",): review_count}, rng, sentences, passages, stratified=False
        )
        rng.shuffle(chosen)
        stats[template] = {
            "auto_pass": count,
            "needs_review": review_count,
            "rows": len(chosen),
            "auto_pool_shortfall": auto_count - count,
            "auto_strata": [
                {
                    "source_kind": k[0],
                    "difficulty": k[1],
                    "pool_size": sizes[k],
                    "sample_size": quotas[k],
                }
                for k in sorted(quotas)
            ],
        }
        for row, item in enumerate(chosen, 1):
            record = {k: v for k, v in item.items() if k != "status"}
            record.update({"pool_status": item["status"], "review_row": row})
            batch.append(record)
    return batch, stats


def highlighted(sentence, spans):
    result, previous = [], 0
    for span in sorted(spans, key=lambda s: (s["start"], s["end"])):
        start, end = span["start"], span["end"]
        if not (previous <= start < end <= len(sentence)):
            raise ValueError("Invalid or overlapping display span")
        result.extend((sentence[previous:start], "【", sentence[start:end], "】"))
        previous = end
    result.append(sentence[previous:])
    return "".join(result)


def teacher_row(item, node_names):
    sentence = item["source"]["sentence"]
    option = item["answer"]["option"]
    display_sentence = highlighted(sentence, [item["target"]]) if item["target"] else sentence
    answer = (
        option
        if item["answer"]["kind"] == "option"
        else highlighted(sentence, item["answer"]["spans"])
    )
    return dict(
        zip(
            HEADERS,
            (
                item["review_row"],
                item["item_id"],
                f"{item['node_id']} {node_names[item['node_id']]}",
                item["prompt_zh"],
                display_sentence,
                "\n".join(item["options"] or []),
                answer,
                "",
                "",
                "",
            ),
            strict=True,
        )
    )


def instruction_rows(total):
    return [
        ("语法候选题抽查", "请逐题核对题干、句子与机器答案，在“老师判定”列选择一项。"),
        ("正确", "题目清楚，机器答案对。"),
        ("答案错", "机器答案不对，请在“正确答案应为”里写出正确答案。"),
        ("题目有歧义", "不止一个合理答案，或者题干说不清。"),
        ("句子不适合出题", "原句残缺、太偏，或者和这个知识点无关。"),
        ("标记说明", "【】标出分类题的目标，或点选题、划边界题的机器答案；原句不改写。"),
        ("预计用时", f"共 {total} 题，按每题 30 秒估算，约 {total / 2:g} 分钟。"),
        ("填写完成后", "请将完整文件交回本项目负责人（发送此表的老师），按题目编号汇总复核。"),
        (
            "后续入库标准",
            "某模板自动通过池的 50 道抽查题中，“答案错”加“题目有歧义”不超过 1 道，"
            "该模板才可以入库；否则调整规则、重新生成、重新抽样。此表不自动执行入库判断。",
        ),
        ("题目来源", "项目 Reading Study 雅思真题标注；题目编号用于与抽查清单的机器可读版对应。"),
    ]


def row_height(values, widths):
    lines = []
    for value, width in zip(values, widths, strict=True):
        parts = str(value).splitlines() or [""]
        lines.append(
            sum(
                max(
                    1,
                    ceil(
                        sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in part)
                        / (width - 3)
                    ),
                )
                for part in parts
            )
        )
    return max(36, max(lines, default=1) * 16 + 12)


def style_sheet(sheet, widths, header=True):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet:
        values = [cell.value or "" for cell in row]
        sheet.row_dimensions[row[0].row].height = row_height(values, widths)
        for cell in row:
            cell.font = Font(name="Arial", size=11, color="202B3A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if header and cell.row == 1:
                cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="30445C")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = Border(right=Side(style="thin", color="FFFFFF"))
            elif header and cell.column >= 8:
                cell.fill = PatternFill("solid", fgColor="FFF4D9")
            elif header and cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F6FA")
    if header:
        sheet.freeze_panes = "D2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.row_dimensions[1].height = 42
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def stable_xlsx(workbook, path):
    """Remove run-dependent core/ZIP timestamps; STORE avoids zlib-version drift."""
    fixed = datetime(2000, 1, 1)
    workbook.properties.created = workbook.properties.modified = fixed
    workbook.properties.creator = "StudyTracker"
    workbook.properties.lastModifiedBy = "StudyTracker"
    memory = BytesIO()
    workbook.save(memory)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(memory) as source, ZipFile(path, "w", compression=ZIP_STORED) as target:
        for name in sorted(source.namelist()):
            payload = source.read(name)
            if name == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+",
                    rb"\g<1>2000-01-01T00:00:00Z",
                    payload,
                )
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.create_system, info.external_attr = 3, 0o600 << 16
            target.writestr(info, payload)


def write_workbook(batch, node_names, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    intro = workbook.active
    intro.title = "说明"
    for row in instruction_rows(len(batch)):
        intro.append(row)
    style_sheet(intro, (24, 105), header=False)
    intro["A1"].font = Font(name="Arial", size=15, bold=True, color="30445C")
    for template in TEMPLATES:
        sheet = workbook.create_sheet(SHEET_NAMES[template])
        sheet.append(HEADERS)
        for item in batch:
            if item["template"] == template:
                sheet.append(list(teacher_row(item, node_names).values()))
        style_sheet(sheet, WIDTHS)
        validation = DataValidation(
            type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True
        )
        validation.showErrorMessage = True
        validation.errorTitle, validation.error = "请选择判定", "请从四个选项中选择。"
        validation.showDropDown = False
        sheet.add_data_validation(validation)
        if sheet.max_row >= 2:
            validation.add(f"H2:H{sheet.max_row}")
    stable_xlsx(workbook, path)
