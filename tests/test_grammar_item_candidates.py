"""Lightweight contract tests: handmade dependency trees, no spaCy import."""

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from collections import Counter
from dataclasses import replace
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from scripts.generate_grammar_candidates import collect
from scripts.grammar_items import spans as roles
from scripts.grammar_items.checks import Reason as R
from scripts.grammar_items.checks import (
    corpus_vocabulary,
    glued_sentence_boundary,
    list_letters,
    multi_sentence_unit,
    review_reasons,
    sentence_reasons,
    suspected_split_word,
    unbalanced,
)
from scripts.grammar_items.parse_view import ParseView, Token, finite, trim_tokens
from scripts.grammar_items.reporting import (
    independent_hits,
    relative_adverb_stats,
    resplit_stats,
    sentence_review_stats,
    template_counts,
)
from scripts.grammar_items.review_batch import (
    HEADERS,
    SHEET_NAMES,
    VERDICTS,
    constrained_sample,
    highlighted,
    largest_remainder,
    sample_batch,
    teacher_row,
    write_workbook,
)
from scripts.grammar_items.spans import classification_kind, locate_spans
from scripts.grammar_items.templates import (
    TEMPLATES,
    close_quotes,
    generate_templates,
    item_id,
    single_quote_role,
)

ROOT = Path(__file__).resolve().parents[1]


def tree(rows):
    """Rows: text, dep, head, pos, tag; offsets are in the returned source."""
    sentence = " ".join(row[0] for row in rows)
    tokens, cursor = [], 0
    for index, (text, dep, head, pos, tag, *forms) in enumerate(rows):
        tokens.append(
            Token(
                index,
                text,
                cursor,
                cursor + len(text),
                pos,
                tag,
                dep,
                head,
                tuple(forms),
                pos == "PUNCT",
            )
        )
        cursor += len(text) + 1
    return sentence, ParseView(tuple(tokens))


def source(sentence, passage="test_p1", sid="A-01"):
    return {
        "source_kind": "reading_test",
        "test_id": "test",
        "passage_id": passage,
        "sentence_id": sid,
        "sentence": sentence,
    }


def generate(sentence, view, tags):
    spans = locate_spans(
        sentence, [{"text": text, "role": role, "level": level} for text, role, level in tags]
    )
    return generate_templates(source(sentence), spans, view)


def by_template(items, template):
    return [item for item in items if item["template"] == template]


def main_tree(marker="when"):
    return tree(
        [
            ("The", "det", 1, "DET", "DT"),
            ("plans", "nsubj", 4, "NOUN", "NNS"),
            ("may", "aux", 4, "AUX", "MD"),
            ("not", "neg", 4, "PART", "RB"),
            ("work", "ROOT", 4, "VERB", "VB"),
            ("out", "prt", 4, "ADP", "RP"),
            (marker, "mark", 8, "SCONJ", "IN"),
            ("prices", "nsubj", 8, "NOUN", "NNS"),
            ("rise", "advcl", 4, "VERB", "VBP"),
            ("today", "npadvmod", 8, "NOUN", "NN"),
            (".", "punct", 4, "PUNCT", "."),
        ]
    )


def main_tags(marker="when", role="time_clause"):
    return [
        ("The plans", "subject", 1),
        ("may not work out", "predicate", 1),
        (f"{marker} prices rise today", role, 2),
    ]


def sampling_pool():
    result = []
    for template in TEMPLATES:
        for status, count in (("auto_pass", 120), ("needs_review", 30)):
            for index in range(count):
                sentence = "The old plans may still work ."
                src = source(sentence, f"{status}_p{index // 2}", f"A-{index}")
                src["source_kind"] = "reading_test" if index % 2 else "reading_jijing"
                start, end = sentence.index("work"), sentence.index("work") + 4
                result.append(
                    {
                        "item_id": item_id(src, template, start, end, "mock"),
                        "generator_version": "grammar-items-v0",
                        "template": template,
                        "node_id": "G-S01",
                        "source": src,
                        "prompt_zh": "点出主句谓语。",
                        "target": None,
                        "options": None,
                        "answer": {
                            "kind": "tokens",
                            "spans": [{"start": start, "end": end, "text": "work"}],
                            "option": None,
                        },
                        "grading": {"mode": "contains_root_subset_of_verb_unit"},
                        "distractors": [],
                        "difficulty": ("easy", "medium", "hard")[index % 3],
                        "status": status,
                        "reasons": [] if status == "auto_pass" else [R.SUBJECT_MISMATCH],
                        "evidence": {},
                    }
                )
    return result


class SpanAndFilterTests(unittest.TestCase):
    def test_ordered_repeated_location(self):
        spans = locate_spans(
            "red and red",
            [
                {"text": "red", "role": "subject", "level": 1},
                {"text": "and", "role": "conjunction", "level": 1},
                {"text": "red", "role": "object", "level": 1},
            ],
        )
        self.assertEqual([s.start for s in spans], [0, 4, 8])
        self.assertTrue(spans[0].ambiguous and spans[2].ambiguous)

    def test_missing_span_is_rejected_without_fallback(self):
        sentence, view = main_tree()
        spans = locate_spans(sentence, [{"text": "absent", "role": "subject", "level": 1}])
        self.assertIn(R.SPAN_NOT_FOUND, sentence_reasons(sentence, spans, view))
        reversed_spans = locate_spans(
            sentence,
            [
                {"text": "prices", "role": "subject", "level": 2},
                {"text": "plans", "role": "subject", "level": 1},
            ],
        )
        self.assertEqual(reversed_spans[1].start, -1)

    def test_role_whitelists_exactly_match_plan(self):
        plan = (ROOT / "docs/GRAMMAR_ITEM_CANDIDATES_PLAN.md").read_text()
        checked = 0
        for line in plan.splitlines():
            parts = [x.strip() for x in line.split("|")]
            if len(parts) != 4:
                continue
            name = parts[1].strip("`")
            if name.endswith("_ROLES") and name != "ADVERBIAL_CLAUSE_ROLES":
                self.assertEqual(getattr(roles, name), set(parts[2].split(", ")))
                checked += 1
        self.assertEqual(checked, 6)
        self.assertEqual(
            roles.ADVERBIAL_CLAUSE_ROLES, {"adverbial_clause", *roles.RELATION_BY_ROLE}
        )

    def test_relation_tables_match_plan(self):
        plan = (ROOT / "docs/GRAMMAR_ITEM_CANDIDATES_PLAN.md").read_text()
        section = plan.split("`RELATION_BY_ROLE`：", 1)[1]
        tables = section.split("`SUBORDINATOR_RELATIONS`（无歧义）：", 1)
        for text, actual in zip(
            tables, (roles.RELATION_BY_ROLE, roles.SUBORDINATOR_RELATIONS), strict=True
        ):
            expected = {}
            for line in text.split("有歧义：", 1)[0].splitlines():
                parts = [x.strip() for x in line.split("|")]
                if len(parts) == 4 and parts[1] in {
                    "条件",
                    "时间",
                    "目的",
                    "让步",
                    "结果",
                    "对比",
                    "原因",
                }:
                    expected.update({role: parts[1] for role in parts[2].split(", ")})
            self.assertEqual(actual, expected)
        self.assertEqual(roles.AMBIGUOUS_SUBORDINATORS, {"while", "since", "as", "so"})

    def test_catch_all_clause_roles_are_never_targets(self):
        excluded = [
            "imperative_clause",
            "correlative_main_clause",
            "clause",
            "explanatory_clause",
            "nominal_clause",
            "interrogative_clause",
            "infinitive_clause",
            "elliptical_clause",
            "passive_clause",
            "complement_clause",
            "reduced_relative_clause",
            "fused_relative_clause",
            "cleft_relative_clause",
            "coordinated_relative_clause",
            "time_condition_clause",
        ]
        for role in excluded:
            with self.subTest(role=role):
                span = locate_spans("a clause", [{"text": "a clause", "role": role, "level": 2}])[0]
                self.assertIsNone(classification_kind(span))

    def test_all_whole_sentence_reasons_are_preserved(self):
        sentence, view = main_tree()
        spans = locate_spans(
            sentence,
            [
                {"text": "The plans", "role": "noun_phrase_fragment", "level": 1},
                {"text": "absent", "role": "purpose_modifier", "level": 2},
            ],
        )
        self.assertEqual(
            set(sentence_reasons(sentence, spans, view)),
            {R.TAGGED_FRAGMENT, R.SPAN_NOT_FOUND, R.UNKNOWN_SPAN},
        )
        self.assertIn(R.PARSE_FAILURE, sentence_reasons(sentence, spans, None))

    def test_clause_attached_to_cc_and_normal_nested_coordinate(self):
        sentence, view = tree(
            [
                ("It", "nsubj", 1, "PRON", "PRP"),
                ("stopped", "ROOT", 1, "VERB", "VBD"),
                ("and", "cc", 1, "CCONJ", "CC"),
                ("when", "mark", 5, "SCONJ", "WRB"),
                ("rain", "nsubj", 5, "NOUN", "NN"),
                ("came", "advcl", 2, "VERB", "VBD"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        self.assertIn(R.CLAUSE_ATTACHED_TO_CONJUNCTION, sentence_reasons(sentence, (), view))
        normal, normal_view = tree(
            [
                ("It", "nsubj", 1, "PRON", "PRP"),
                ("lasts", "ROOT", 1, "VERB", "VBZ"),
                ("and", "cc", 1, "CCONJ", "CC"),
                ("when", "mark", 6, "SCONJ", "WRB"),
                ("it", "nsubj", 6, "PRON", "PRP"),
                ("does", "aux", 6, "AUX", "VBZ"),
                ("break", "advcl", 8, "VERB", "VB"),
                ("it", "nsubj", 8, "PRON", "PRP"),
                ("shatters", "conj", 1, "VERB", "VBZ"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        self.assertNotIn(
            R.CLAUSE_ATTACHED_TO_CONJUNCTION, sentence_reasons(normal, (), normal_view)
        )

    def test_quotes_brackets_and_apostrophes(self):
        for text in [
            "It's fine.",
            "The students' books.",
            "‘It's fine,’ she said.",
            'She said "fine" (twice).',
        ]:
            self.assertFalse(unbalanced(text), text)
        for text in ["Be careful of little lives.'", "(unfinished", "“unfinished", "(wrong]"]:
            self.assertTrue(unbalanced(text), text)

    def test_filter_word_lengths_root_heading_and_level_one(self):
        sentence, view = main_tree()
        short = ParseView(view.tokens[:5])
        self.assertIn(R.TOO_SHORT, sentence_reasons(sentence, (), short))
        long = ParseView(tuple(replace(view.tokens[0], i=i) for i in range(61)))
        self.assertIn(R.TOO_LONG, sentence_reasons(sentence, (), long))
        wrong_root = ParseView(
            tuple(replace(t, pos="NOUN") if t.dep == "ROOT" else t for t in view.tokens)
        )
        self.assertIn(R.FRAGMENT, sentence_reasons(sentence, (), wrong_root))
        for role, reason in [
            ("heading", R.HEADING),
            ("relative_clause", R.LEVEL_ONE_SUBORDINATE),
            ("noun_clause", R.LEVEL_ONE_SUBORDINATE),
        ]:
            spans = locate_spans(sentence, [{"text": "The plans", "role": role, "level": 1}])
            self.assertIn(reason, sentence_reasons(sentence, spans, view))


class TemplateTests(unittest.TestCase):
    def test_t1_verb_unit_root_distractor_and_grading(self):
        sentence, view = main_tree()
        candidates, _ = generate(sentence, view, main_tags())
        item = by_template(candidates, TEMPLATES[0])[0]
        self.assertEqual(item["status"], "auto_pass")
        self.assertEqual(
            [s["text"] for s in item["answer"]["spans"]], ["may", "not", "work", "out"]
        )
        self.assertEqual([s["text"] for s in item["distractors"]], ["rise"])
        self.assertEqual(item["distractors"][0]["roles"], ["time_clause"])
        self.assertEqual(item["node_id"], "G-L01")
        self.assertEqual(item["grading"]["required_root"]["text"], "work")
        self.assertIsNone(item["target"])

    def test_t1_wrong_root_inside_relative_clause_needs_review(self):
        sentence, view = tree(
            [
                ("The", "det", 1, "DET", "DT"),
                ("people", "nsubj", 3, "NOUN", "NNS"),
                ("who", "nsubj", 3, "PRON", "WP"),
                ("arrived", "ROOT", 3, "VERB", "VBD"),
                ("today", "npadvmod", 3, "NOUN", "NN"),
                ("completed", "xcomp", 3, "VERB", "VBD"),
                ("the", "det", 7, "DET", "DT"),
                ("task", "dobj", 5, "NOUN", "NN"),
                (".", "punct", 3, "PUNCT", "."),
            ]
        )
        candidates, _ = generate(
            sentence,
            view,
            [
                ("The people", "subject", 1),
                ("who arrived today", "relative_clause", 2),
                ("completed", "verb", 1),
            ],
        )
        item = by_template(candidates, TEMPLATES[0])[0]
        self.assertEqual(item["status"], "needs_review")
        self.assertIn(R.PARSER_ROOT_OUTSIDE_VERB_SPAN, item["reasons"])
        self.assertIn(R.PARSER_ROOT_INSIDE_SUBORDINATE_SPAN, item["reasons"])

    def test_t1_rejects_parser_coordinated_verbs(self):
        sentence, view = main_tree()
        view = ParseView(
            tuple(replace(t, dep="conj") if t.text == "rise" else t for t in view.tokens)
        )
        candidates, rejects = generate(sentence, view, main_tags())
        self.assertFalse(by_template(candidates, TEMPLATES[0]))
        self.assertIn(R.COORDINATED_MAIN_VERBS, by_template(rejects, TEMPLATES[0])[0]["reasons"])

    def test_main_selection_rejections(self):
        sentence, view = main_tree()
        cases = [
            ([], R.NO_MAIN_VERB_TAG),
            ([("may", "verb", 1), ("work out", "verb", 1)], R.MULTIPLE_MAIN_VERBS),
            ([("may not work out", "coordinated_predicate", 1)], R.COORDINATED_MAIN_VERBS),
            (
                [
                    ("may not work out", "predicate", 1),
                    ("when prices rise today", "coordinate_clause", 1),
                ],
                R.COORDINATED_MAIN_CLAUSES,
            ),
        ]
        for tags, reason in cases:
            _, rejected = generate(sentence, view, tags)
            self.assertIn(reason, by_template(rejected, TEMPLATES[0])[0]["reasons"])

    def test_t2_exact_subtree_and_boundary_mismatch(self):
        sentence, view = main_tree()
        candidates, _ = generate(sentence, view, main_tags())
        item = by_template(candidates, TEMPLATES[1])[0]
        self.assertEqual(item["status"], "auto_pass")
        self.assertEqual(item["answer"]["spans"][0]["text"], "when prices rise today")
        tags = main_tags()[:-1] + [("when prices rise", "time_clause", 2)]
        candidates, _ = generate(sentence, view, tags)
        self.assertIn(R.BOUNDARY_MISMATCH, by_template(candidates, TEMPLATES[1])[0]["reasons"])

    def test_t2_coordinate_leading_conjunction(self):
        sentence, view = tree(
            [
                ("The", "det", 1, "DET", "DT"),
                ("team", "nsubj", 2, "NOUN", "NN"),
                ("worked", "ROOT", 2, "VERB", "VBD"),
                ("hard", "advmod", 2, "ADV", "RB"),
                ("and", "cc", 2, "CCONJ", "CC"),
                ("the", "det", 6, "DET", "DT"),
                ("students", "nsubj", 7, "NOUN", "NNS"),
                ("rested", "conj", 2, "VERB", "VBD"),
                ("quietly", "advmod", 7, "ADV", "RB"),
                (".", "punct", 2, "PUNCT", "."),
            ]
        )
        candidates, _ = generate(
            sentence, view, [("and the students rested quietly", "coordinate_clause", 1)]
        )
        item = by_template(candidates, TEMPLATES[1])[0]
        self.assertEqual(item["status"], "auto_pass")
        self.assertEqual(item["answer"]["spans"][0]["text"], "the students rested quietly")
        self.assertTrue(item["grading"]["allow_leading_conjunction"])

    def test_t2_multiple_same_type_does_not_block_t3(self):
        sentence, view = main_tree()
        tags = main_tags() + [("rise today", "time_clause", 3)]
        candidates, rejections = generate(sentence, view, tags)
        self.assertFalse(by_template(candidates, TEMPLATES[1]))
        self.assertEqual(len(by_template(candidates, TEMPLATES[2])), 2)
        self.assertTrue(
            all(
                R.MULTIPLE_TARGETS_SAME_TYPE in r["reasons"]
                for r in by_template(rejections, TEMPLATES[1])
            )
        )

    def test_t3_requires_finite_clause_but_t4_checks_relation_only(self):
        sentence, view = main_tree("although")
        view = ParseView(
            tuple(replace(t, tag="VBG") if t.text == "rise" else t for t in view.tokens)
        )
        candidates, _ = generate(sentence, view, main_tags("although", "concessive_clause"))
        self.assertIn(R.CLAUSE_TYPE_MISMATCH, by_template(candidates, TEMPLATES[2])[0]["reasons"])
        self.assertEqual(by_template(candidates, TEMPLATES[3])[0]["status"], "auto_pass")

    def test_t4_concessive_ambiguous_and_mismatch(self):
        for marker, role, reason in [
            ("although", "concessive_clause", None),
            ("while", "time_clause", R.SUBORDINATOR_AMBIGUOUS),
            ("because", "time_clause", R.SUBORDINATOR_MISMATCH),
            ("wherever", "time_clause", R.SUBORDINATOR_UNKNOWN),
        ]:
            with self.subTest(marker=marker):
                sentence, view = main_tree(marker)
                candidates, _ = generate(sentence, view, main_tags(marker, role))
                item = by_template(candidates, TEMPLATES[3])[0]
                self.assertEqual(item["status"], "needs_review" if reason else "auto_pass")
                if reason:
                    self.assertIn(reason, item["reasons"])

    def test_t4_multiword_subordinator_not_misclassified_as_ambiguous(self):
        from scripts.grammar_items.templates import Context, subordinator

        for phrase in ["so that", "even though", "as long as", "as soon as", "in order that"]:
            words = phrase.split()
            rows = [
                (word, "mark" if i == len(words) - 1 else "advmod", len(words) + 1, "SCONJ", "IN")
                for i, word in enumerate(words)
            ]
            rows += [
                ("prices", "nsubj", len(words) + 1, "NOUN", "NNS"),
                ("rise", "ROOT", len(words) + 1, "VERB", "VBP"),
            ]
            sentence, view = tree(rows)
            ctx = Context(source(sentence), (), view)
            self.assertEqual(subordinator(ctx, view.tokens, view.root), phrase)

    def test_t4_plain_adverbial_has_no_relation_label(self):
        sentence, view = main_tree()
        _, rejects = generate(sentence, view, main_tags(role="adverbial_clause"))
        self.assertEqual(by_template(rejects, TEMPLATES[3])[0]["reasons"], [R.NO_RELATION_LABEL])

    def test_t5_bad_subject_does_not_propagate_to_t1_or_t6(self):
        sentence, view = tree(
            [
                ("By", "prep", 8, "ADP", "IN"),
                ("the", "det", 3, "DET", "DT"),
                ("late", "amod", 3, "ADJ", "JJ"),
                ("1980s", "pobj", 0, "NOUN", "NNS"),
                ("the", "det", 6, "DET", "DT"),
                ("Modern", "compound", 6, "PROPN", "NNP"),
                ("Movement", "nsubj", 8, "PROPN", "NNP"),
                ("had", "aux", 8, "AUX", "VBD"),
                ("faded", "ROOT", 8, "VERB", "VBN"),
                (".", "punct", 8, "PUNCT", "."),
            ]
        )
        candidates, _ = generate(
            sentence,
            view,
            [
                ("By the late 1980s the Modern Movement", "subject", 1),
                ("had faded", "predicate", 1),
            ],
        )
        subject = by_template(candidates, TEMPLATES[4])[0]
        self.assertEqual(subject["status"], "needs_review")
        self.assertIn(R.SUBJECT_MISMATCH, subject["reasons"])
        self.assertEqual(by_template(candidates, TEMPLATES[0])[0]["status"], "auto_pass")
        self.assertEqual(by_template(candidates, TEMPLATES[0])[0]["node_id"], "G-S01")
        self.assertEqual(by_template(candidates, TEMPLATES[5])[0]["status"], "auto_pass")
        self.assertEqual(by_template(candidates, TEMPLATES[5])[0]["node_id"], "G-N01")

    def test_t5_existential_split_and_subject_after_verb(self):
        sentence, view = main_tree()
        cases = [
            (main_tags()[1:], view, R.SPLIT_SUBJECT),
            (
                [("may not work out", "predicate", 1), ("prices", "subject", 1)],
                view,
                R.SUBJECT_AFTER_VERB,
            ),
            (
                main_tags(),
                ParseView(tuple(replace(t, dep="expl") if t.i == 1 else t for t in view.tokens)),
                R.EXISTENTIAL,
            ),
        ]
        for tags, parsed, reason in cases:
            _, rejected = generate(sentence, parsed, tags)
            self.assertIn(reason, by_template(rejected, TEMPLATES[4])[0]["reasons"])

    def test_finiteness_rule_and_nonfinite_node_mapping(self):
        sentence, view = main_tree()
        self.assertTrue(finite(view.root, view))  # Finite modal child.
        vbn = replace(view.root, tag="VBN", morph_verbform=("Part",))
        empty_children = ParseView((replace(vbn, i=0, head=0),))
        self.assertFalse(finite(empty_children.root, empty_children))
        self.assertTrue(finite(replace(vbn, morph_verbform=("Fin",)), empty_children))
        for role, node in [("participial_phrase", "G-N04"), ("gerund_phrase", "G-N01")]:
            parsed = ParseView(
                tuple(
                    replace(t, tag="VBG", dep="acl") if t.text == "rise" else t for t in view.tokens
                )
            )
            candidates, _ = generate(
                sentence, parsed, main_tags()[:-1] + [("prices rise today", role, 2)]
            )
            nonfinite = [
                i
                for i in by_template(candidates, TEMPLATES[5])
                if i["evidence"]["variant"] == "non_finite"
            ][0]
            self.assertEqual(nonfinite["node_id"], node)

    def test_ids_offsets_field_order_and_repeatability(self):
        sentence, view = main_tree()
        first, _ = generate(sentence, view, main_tags())
        second, _ = generate(sentence, view, main_tags())
        self.assertEqual(first, second)
        self.assertEqual(len({i["item_id"] for i in first}), len(first))
        expected = [
            "item_id",
            "generator_version",
            "template",
            "node_id",
            "source",
            "prompt_zh",
            "target",
            "options",
            "answer",
            "grading",
            "distractors",
            "difficulty",
            "status",
            "reasons",
            "evidence",
        ]
        for item in first:
            self.assertEqual(list(item), expected)
            for span in [
                *item["answer"]["spans"],
                *item["distractors"],
                *([item["target"]] if item["target"] else []),
            ]:
                self.assertEqual(sentence[span["start"] : span["end"]], span["text"])
        self.assertNotEqual(
            item_id(source(sentence), "a", 0, 1, "x"), item_id(source(sentence), "a", 0, 1, "y")
        )

    def test_ambiguous_target_text_cannot_auto_pass(self):
        sentence, view = main_tree()
        tags = main_tags()
        spans = locate_spans(
            sentence, [{"text": t, "role": r, "level": level} for t, r, level in tags]
        )
        spans = (*spans[:2], replace(spans[2], ambiguous=True))
        candidates, _ = generate_templates(source(sentence), spans, view)
        item = by_template(candidates, TEMPLATES[1])[0]
        self.assertEqual(item["status"], "needs_review")
        self.assertIn(R.AMBIGUOUS_SPAN_LOCATION, item["reasons"])


class SamplingAndOutputTests(unittest.TestCase):
    def test_largest_remainder_exact_totals_and_ties(self):
        self.assertEqual(
            largest_remainder({("a",): 5, ("b",): 3, ("c",): 2}, 7),
            {("a",): 4, ("b",): 2, ("c",): 1},
        )
        self.assertEqual(largest_remainder({("b",): 1, ("a",): 1}, 1), {("a",): 1, ("b",): 0})

    def test_seed_strata_caps_and_blind_teacher_rows(self):
        pool = sampling_pool()
        batch, stats = sample_batch(pool, 20261018)
        self.assertEqual((batch, stats), sample_batch(pool, 20261018))
        self.assertNotEqual(batch, sample_batch(pool, 20261019)[0])
        self.assertEqual(len(batch), 360)
        for template in TEMPLATES:
            items = by_template(batch, template)
            self.assertEqual(
                Counter(i["pool_status"] for i in items), {"auto_pass": 50, "needs_review": 10}
            )
            keys = [(i["source"]["passage_id"], i["source"]["sentence_id"]) for i in items]
            self.assertEqual(len(keys), len(set(keys)))
            self.assertLessEqual(max(Counter(k[0] for k in keys).values()), 2)
            for stratum in stats[template]["auto_strata"]:
                self.assertEqual(
                    sum(
                        i["pool_status"] == "auto_pass"
                        and i["source"]["source_kind"] == stratum["source_kind"]
                        and i["difficulty"] == stratum["difficulty"]
                        for i in items
                    ),
                    stratum["sample_size"],
                )
        for item in batch:
            self.assertNotIn("status", item)
            row = teacher_row(item, {"G-S01": "识别谓语动词单位"})
            self.assertEqual(tuple(row), HEADERS)
            serialized = json.dumps(row, ensure_ascii=False)
            self.assertNotIn("auto_pass", serialized)
            self.assertNotIn("needs_review", serialized)
            self.assertNotIn("status", serialized)
            self.assertIn("【work】", row["机器答案"])

    def test_shortfall_and_impossible_sample_are_explicit(self):
        pool = sampling_pool()
        short = [
            i for i in pool if i["status"] != "auto_pass" or int(i["source"]["sentence_id"][2:]) < 8
        ]
        batch, stats = sample_batch(short, 1)
        self.assertEqual(len(batch), 6 * 18)
        self.assertTrue(all(s["auto_pool_shortfall"] == 42 for s in stats.values()))
        with self.assertRaisesRegex(ValueError, "review slots"):
            sample_batch([i for i in pool if i["status"] == "auto_pass"], 1)

    def test_display_brackets_respect_original_offsets(self):
        sentence = "They had lost out."
        self.assertEqual(
            highlighted(sentence, [{"start": 5, "end": 8}, {"start": 9, "end": 13}]),
            "They 【had lost】 out.",
        )
        self.assertEqual(
            highlighted(sentence, [{"start": 5, "end": 8}, {"start": 14, "end": 17}]),
            "They 【had】 lost 【out】.",
        )
        with self.assertRaises(ValueError):
            highlighted(sentence, [{"start": 5, "end": 12}, {"start": 9, "end": 13}])

    def test_independent_filter_counts_not_short_circuited(self):
        src = source("A fragment has overlapping rejection reasons .")
        report = independent_hits(
            [(src, [R.UNKNOWN_SPAN, R.TAGGED_FRAGMENT, R.CLAUSE_ATTACHED_TO_CONJUNCTION])]
        )
        self.assertEqual(report[R.TAGGED_FRAGMENT]["actual"], 1)
        self.assertEqual(report[R.CLAUSE_ATTACHED_TO_CONJUNCTION]["actual"], 1)
        counts = template_counts([], [{"scope": "sentence", "template": "all"}])
        self.assertTrue(all(c["rejected"] == 1 for c in counts.values()))

    def test_imports_work_with_spacy_blocked(self):
        code = """
import sys
class BlockSpacy:
    def find_spec(self, fullname, path=None, target=None):
        if fullname == "spacy" or fullname.startswith("spacy."):
            raise AssertionError("spaCy imported by lightweight logic")
sys.meta_path.insert(0, BlockSpacy())
import scripts.generate_grammar_candidates
import scripts.grammar_items.templates
import scripts.grammar_items.review_batch
"""
        result = subprocess.run(
            [sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True
        )
        self.assertEqual(result.returncode, 0, result.stderr)

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl is optional in CI")
    def test_xlsx_blindness_dropdowns_and_byte_reproducibility(self):
        from openpyxl import load_workbook

        batch, _ = sample_batch(sampling_pool(), 20261018)
        with tempfile.TemporaryDirectory() as directory:
            a, b = Path(directory) / "a.xlsx", Path(directory) / "b.xlsx"
            write_workbook(batch, {"G-S01": "识别谓语动词单位"}, a)
            write_workbook(batch, {"G-S01": "识别谓语动词单位"}, b)
            self.assertEqual(sha256(a.read_bytes()).hexdigest(), sha256(b.read_bytes()).hexdigest())
            workbook = load_workbook(a)
            self.assertEqual(workbook.sheetnames, ["说明", *SHEET_NAMES.values()])
            for sheet in list(workbook)[1:]:
                self.assertEqual((sheet.max_row, sheet.max_column), (61, 10))
                self.assertEqual(sheet.freeze_panes, "D2")
                dv = sheet.data_validations.dataValidation[0]
                self.assertEqual(dv.formula1, '"' + ",".join(VERDICTS) + '"')
                self.assertEqual(str(dv.sqref), "H2:H61")
                self.assertFalse(dv.showDropDown)
                self.assertTrue(sheet["E2"].alignment.wrap_text)
                for row in sheet:
                    for cell in row:
                        self.assertIsNone(cell.comment)
                        self.assertNotIn("auto_pass", str(cell.value))
                        self.assertNotIn("needs_review", str(cell.value))
            with ZipFile(a) as archive:
                self.assertNotIn(
                    "auto_pass",
                    "".join(
                        archive.read(n).decode() for n in archive.namelist() if n.endswith(".xml")
                    ),
                )


class V02Tests(unittest.TestCase):
    """Plan v0.2: new sentence filters, relative-adverb risk and full verb units."""

    def test_close_quotes_reattaches_split_quote(self):
        sentence = 'they called "loss aversion." For "apprentice" programmes'
        end = sentence.index("aversion") + len("aversion")
        span = {"start": 0, "end": end, "text": sentence[:end]}
        fixed = close_quotes(sentence, span, 0, sentence.index(" For"))
        self.assertEqual(fixed["text"], 'they called "loss aversion."')
        start = sentence.index("apprentice")
        span = {"start": start, "end": len(sentence), "text": sentence[start:]}
        fixed = close_quotes(sentence, span, sentence.index('"apprentice'), len(sentence))
        self.assertEqual(fixed["text"], '"apprentice" programmes')
        sentence2 = "someone who 'became rich faster than ever before'."
        start2 = sentence2.index("who")
        end2 = sentence2.index("before") + len("before")
        span2 = {"start": start2, "end": end2, "text": sentence2[start2:end2]}
        fixed = close_quotes(sentence2, span2, start2, len(sentence2))
        self.assertEqual(fixed["text"], "who 'became rich faster than ever before'")
        owners = "the students' books were lost"
        span3 = {"start": 4, "end": 20, "text": owners[4:20]}
        self.assertIs(close_quotes(owners, span3, 4, 20), span3)
        balanced = {"start": 0, "end": 4, "text": "they"}
        self.assertIs(close_quotes(sentence, balanced, 0, len(sentence)), balanced)

    def test_glued_sentence_boundary(self):
        self.assertTrue(glued_sentence_boundary('such as these."Manufacturers are not aware.'))
        self.assertTrue(glued_sentence_boundary("the future of Asia.Thailand's boom"))
        self.assertFalse(glued_sentence_boundary('such as these." Manufacturers are not aware.'))
        self.assertFalse(glued_sentence_boundary("J.K.Rowling wrote the books."))

    def test_multi_sentence_unit(self):
        self.assertTrue(multi_sentence_unit("It failed. They tried again."))
        self.assertFalse(multi_sentence_unit("He met Dr. Smith yesterday."))
        self.assertFalse(multi_sentence_unit("'All music has a message.' he says."))

    def test_multi_sentence_unit_known_non_ends(self):
        # False hits found when the 51 units were re-split (2026-10-19).
        self.assertFalse(multi_sentence_unit("ice can be reflected in Mt. Cook Region."))
        self.assertFalse(
            multi_sentence_unit('requests ("Do" vs. "Don\'t" demand contexts) were explored.')
        )
        self.assertFalse(multi_sentence_unit("1. Valid assessment considers cultural diversity."))
        self.assertFalse(multi_sentence_unit('"What\'s different?" Pepperberg asked.'))
        self.assertFalse(multi_sentence_unit('"Was there someone like him?" Dapena asks.'))
        self.assertTrue(multi_sentence_unit('"Is it?" Perhaps he\'s right.'))
        self.assertTrue(multi_sentence_unit("It rose by 2. Then it fell."))

    def test_new_filters_are_sentence_reasons(self):
        reasons = sentence_reasons("It failed. They tried again.Then it worked.", [], None)
        self.assertIn(R.GLUED_SENTENCE_BOUNDARY, reasons)
        self.assertIn(R.MULTI_SENTENCE_UNIT, reasons)

    def relative_adverb_items(self):
        sentence, view = tree(
            [
                ("People", "nsubj", 1, "NOUN", "NNS"),
                ("save", "ROOT", 1, "VERB", "VBP"),
                ("money", "dobj", 1, "NOUN", "NN"),
                ("when", "mark", 5, "SCONJ", "WRB"),
                ("prices", "nsubj", 5, "NOUN", "NNS"),
                ("rise", "advcl", 1, "VERB", "VBP"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        tags = [
            ("People", "subject", 1),
            ("save", "verb", 1),
            ("money", "object", 1),
            ("when prices rise", "time_clause", 2),
        ]
        return generate(sentence, view, tags)[0]

    def test_relative_adverb_risk_after_noun(self):
        items = self.relative_adverb_items()
        for template in ("clause_boundary", "clause_type_classify", "adverbial_relation"):
            adverbial = [i for i in by_template(items, template) if "when" in str(i)]
            self.assertTrue(adverbial, template)
            for item in adverbial:
                self.assertIn(R.POSSIBLE_RELATIVE_ADVERB, item["reasons"])
                self.assertEqual(item["status"], "needs_review")

    def test_no_relative_adverb_risk_after_verb_particle(self):
        sentence, view = main_tree()
        items, _ = generate(sentence, view, main_tags())
        relation = by_template(items, "adverbial_relation")
        self.assertTrue(relation)
        self.assertNotIn(R.POSSIBLE_RELATIVE_ADVERB, relation[0]["reasons"])

    def test_main_verb_target_is_contiguous_unit(self):
        sentence, view = main_tree()
        items, _ = generate(sentence, view, main_tags())
        main = [
            i for i in by_template(items, "finite_or_not") if i["answer"]["option"] == "主句谓语"
        ]
        self.assertEqual(main[0]["target"]["text"], "may not work out")
        self.assertEqual(len(main[0]["target"]["parts"]), 1)

    def test_discontinuous_unit_has_two_parts(self):
        sentence, view = tree(
            [
                ("The", "det", 1, "DET", "DT"),
                ("team", "nsubj", 4, "NOUN", "NN"),
                ("has", "aux", 4, "AUX", "VBZ"),
                ("also", "advmod", 4, "ADV", "RB"),
                ("identified", "ROOT", 4, "VERB", "VBN"),
                ("problems", "dobj", 4, "NOUN", "NNS"),
                (".", "punct", 4, "PUNCT", "."),
            ]
        )
        tags = [("The team", "subject", 1), ("has also identified", "verb_phrase", 1)]
        items, _ = generate(sentence, view, [*tags, ("problems", "object", 1)])
        item = by_template(items, "finite_or_not")[0]
        self.assertEqual([p["text"] for p in item["target"]["parts"]], ["has", "identified"])
        row = teacher_row({**item, "review_row": 1}, {item["node_id"]: "名称"})
        self.assertIn("【has】 also 【identified】", row[HEADERS[4]])

    def test_infinitive_target_includes_to(self):
        sentence, view = tree(
            [
                ("She", "nsubj", 1, "PRON", "PRP"),
                ("wants", "ROOT", 1, "VERB", "VBZ"),
                ("to", "aux", 3, "PART", "TO"),
                ("introduce", "xcomp", 1, "VERB", "VB"),
                ("changes", "dobj", 3, "NOUN", "NNS"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        tags = [
            ("She", "subject", 1),
            ("wants", "verb", 1),
            ("to introduce changes", "infinitive_phrase", 2),
        ]
        items, _ = generate(sentence, view, tags)
        non_finite = [
            i for i in by_template(items, "finite_or_not") if i["answer"]["option"] != "主句谓语"
        ]
        self.assertEqual(non_finite[0]["target"]["text"], "to introduce")

    def test_passive_clause_verb_and_distractor_show_was_born(self):
        sentence, view = tree(
            [
                ("That", "nsubj", 1, "PRON", "DT"),
                ("was", "ROOT", 1, "AUX", "VBD"),
                ("the", "det", 3, "DET", "DT"),
                ("moment", "attr", 1, "NOUN", "NN"),
                ("when", "advmod", 7, "SCONJ", "WRB"),
                ("cinema", "nsubjpass", 7, "NOUN", "NN"),
                ("was", "auxpass", 7, "AUX", "VBD"),
                ("born", "relcl", 3, "VERB", "VBN"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        tags = [
            ("That", "subject", 1),
            ("was", "linking_verb", 1),
            ("the moment when cinema was born", "predicative", 1),
            ("when cinema was born", "relative_clause", 2),
        ]
        items, _ = generate(sentence, view, tags)
        clause = [
            i
            for i in by_template(items, "finite_or_not")
            if i["answer"]["option"] == "从句里的谓语"
        ]
        self.assertEqual(clause[0]["target"]["text"], "was born")
        distractors = by_template(items, "main_verb_tap")[0]["distractors"]
        self.assertEqual([d["text"] for d in distractors], ["was born"])


class CodexReviewFixTests(unittest.TestCase):
    """Fixes for the independent v0.2 review (F1, F2, F5, S1)."""

    def test_split_word_uses_corpus_vocabulary(self):
        vocabulary = corpus_vocabulary(["It was calculated carefully.", "They calculated it."])
        self.assertTrue(suspected_split_word("They had calcu lated that it works.", vocabulary))
        self.assertFalse(suspected_split_word("They had calculated that it works.", vocabulary))
        common = corpus_vocabulary(["red and red and red ."])
        self.assertFalse(suspected_split_word("red and red", common))
        reasons = sentence_reasons("They had calcu lated it.", [], None, vocabulary)
        self.assertIn(R.SUSPECTED_SPLIT_WORD, reasons)
        self.assertNotIn(R.SUSPECTED_SPLIT_WORD, sentence_reasons("calcu lated", [], None))

    def test_leading_stray_punctuation(self):
        def reasons(sentence):
            return sentence_reasons(sentence, [], None)

        self.assertIn(R.LEADING_STRAY_PUNCTUATION, reasons(". To encourage reform, he spoke."))
        self.assertIn(R.LEADING_STRAY_PUNCTUATION, reasons(", and then it rained."))
        self.assertNotIn(R.LEADING_STRAY_PUNCTUATION, reasons("... I think this is fine."))
        self.assertNotIn(R.LEADING_STRAY_PUNCTUATION, reasons("To encourage reform, he spoke."))

    def test_control_character(self):
        reasons = sentence_reasons("It cost \x00220 billion in total.", [], None)
        self.assertIn(R.CONTROL_CHARACTER, reasons)
        self.assertNotIn(R.CONTROL_CHARACTER, sentence_reasons("It cost $220 billion.", [], None))

    def interrupted_items(self):
        sentence, view = tree(
            [
                ("They", "nsubj", 6, "PRON", "PRP"),
                ("do", "aux", 6, "AUX", "VBP"),
                ("not", "neg", 6, "PART", "RB"),
                (",", "punct", 6, "PUNCT", ","),
                ("however", "advmod", 6, "ADV", "RB"),
                (",", "punct", 6, "PUNCT", ","),
                ("handle", "ROOT", 6, "VERB", "VB"),
                ("it", "dobj", 6, "PRON", "PRP"),
                (".", "punct", 6, "PUNCT", "."),
            ]
        )
        tags = [("They", "subject", 1), ("handle", "verb", 1), ("it", "object", 1)]
        return generate(sentence, view, tags)[0]

    def test_verb_unit_outside_tag_span_is_shown_and_downgraded(self):
        items = self.interrupted_items()
        t6 = by_template(items, "finite_or_not")[0]
        self.assertEqual([p["text"] for p in t6["target"]["parts"]], ["do not", "handle"])
        self.assertIn(R.VERB_UNIT_OUTSIDE_SPAN, t6["reasons"])
        self.assertEqual(t6["status"], "needs_review")

    def test_t1_answer_unchanged_but_downgraded(self):
        t1 = by_template(self.interrupted_items(), "main_verb_tap")[0]
        self.assertEqual([s["text"] for s in t1["answer"]["spans"]], ["handle"])
        self.assertIn(R.VERB_UNIT_OUTSIDE_SPAN, t1["reasons"])

    def test_inversion_and_particle_outside_tag_span(self):
        cases = {
            "inversion": (
                [
                    ("Only", "advmod", 1, "ADV", "RB"),
                    ("then", "advmod", 4, "ADV", "RB"),
                    ("did", "aux", 4, "AUX", "VBD"),
                    ("they", "nsubj", 4, "PRON", "PRP"),
                    ("take", "ROOT", 4, "VERB", "VB"),
                    ("action", "dobj", 4, "NOUN", "NN"),
                    (".", "punct", 4, "PUNCT", "."),
                ],
                [("they", "subject", 1), ("take", "verb", 1), ("action", "object", 1)],
                ["did", "take"],
            ),
            "particle": (
                [
                    ("People", "nsubj", 2, "NOUN", "NNS"),
                    ("can", "aux", 2, "AUX", "MD"),
                    ("put", "ROOT", 2, "VERB", "VB"),
                    ("it", "dobj", 2, "PRON", "PRP"),
                    ("down", "prt", 2, "ADP", "RP"),
                    ("to", "prep", 2, "ADP", "IN"),
                    ("luck", "pobj", 5, "NOUN", "NN"),
                    (".", "punct", 2, "PUNCT", "."),
                ],
                [("People", "subject", 1), ("can put", "verb_phrase", 1), ("it", "object", 1)],
                ["can put", "down"],
            ),
        }
        for name, (rows, tags, parts) in cases.items():
            with self.subTest(name):
                sentence, view = tree(rows)
                items, _ = generate(sentence, view, tags)
                t6 = by_template(items, "finite_or_not")[0]
                self.assertEqual([p["text"] for p in t6["target"]["parts"]], parts)
                self.assertEqual(t6["status"], "needs_review")
                t1 = by_template(items, "main_verb_tap")[0]
                self.assertIn(R.VERB_UNIT_OUTSIDE_SPAN, t1["reasons"])

    def test_report_lists_relative_adverb_targets_and_resplit_counts(self):
        src = source("People save money when prices rise .")
        item = {
            "item_id": "gi_v02_x",
            "template": "adverbial_relation",
            "source": src,
            "reasons": [R.POSSIBLE_RELATIVE_ADVERB],
            "evidence": {"tag_span": {"start": 18, "end": 34}},
        }
        stats = relative_adverb_stats([item, {**item, "template": "clause_boundary"}])
        self.assertEqual(stats["distinct_targets"], 1)
        self.assertEqual(stats["targets"][0]["target"], "when prices rise")
        self.assertEqual(
            stats["targets"][0]["templates"], ["adverbial_relation", "clause_boundary"]
        )
        rows = [(source("x", sid=sid), []) for sid in ("A-01", "A-02", "A-02b", "A-02c")]
        resplit = resplit_stats(rows, baseline=2)
        self.assertEqual((resplit["sentences_before"], resplit["sentences_after"]), (2, 4))
        self.assertEqual((resplit["added"], resplit["removed"], resplit["net_change"]), (2, 0, 2))
        # A-03 merged into A-02 and A-01 split in two: one removed, one added (review N5).
        rows = [(source("x", sid=sid), []) for sid in ("A-01", "A-01b", "A-02")]
        resplit = resplit_stats(rows, baseline=3)
        self.assertEqual((resplit["added"], resplit["removed"], resplit["net_change"]), (1, 1, 0))

    def test_truncated_clause_target_is_not_auto_pass(self):
        sentence, view = tree(
            [
                ("Prices", "nsubj", 1, "NOUN", "NNS"),
                ("rose", "ROOT", 1, "VERB", "VBD"),
                ("because", "mark", 4, "SCONJ", "IN"),
                ("demand", "nsubj", 4, "NOUN", "NN"),
                ("grew", "advcl", 1, "VERB", "VBD"),
                ("quickly", "advmod", 4, "ADV", "RB"),
                ("abroad", "advmod", 4, "ADV", "RB"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        tags = [
            ("Prices", "subject", 1),
            ("rose", "verb", 1),
            ("because demand grew", "reason_clause", 2),
        ]
        items, _ = generate(sentence, view, tags)
        for template in ("clause_boundary", "clause_type_classify", "adverbial_relation"):
            item = by_template(items, template)[0]
            self.assertIn(R.BOUNDARY_MISMATCH, item["reasons"], template)
            self.assertEqual(item["status"], "needs_review")

    def test_apostrophes_are_not_quote_marks(self):
        self.assertIsNone(single_quote_role("in the '80s", 7))
        self.assertIsNone(single_quote_role("it's", 2))
        self.assertEqual(single_quote_role("students' books", 8), "ambiguous")
        self.assertEqual(single_quote_role("James' car", 5), "ambiguous")
        sentence = "the 'merchant of death', someone"
        self.assertEqual(single_quote_role(sentence, 4), "open")
        self.assertEqual(single_quote_role(sentence, 22), "close")
        quoted = "'The students' books'"
        span = {"start": 1, "end": 20, "text": quoted[1:20]}
        self.assertIs(close_quotes(quoted, span, 0, len(quoted)), span)

    def test_close_quotes_reattaches_bracket_within_raw_span(self):
        sentence = "The stepwell of Rani Ki Vav (Queen's Stepwell) is famous."
        close = sentence.index(")")
        span = {"start": 0, "end": close, "text": sentence[:close]}
        closed = close_quotes(sentence, span, 0, close + 1)
        self.assertEqual(closed["text"], "The stepwell of Rani Ki Vav (Queen's Stepwell)")
        self.assertIs(close_quotes(sentence, span, 0, close), span)

    def test_coordinated_subordinate_clause_is_not_coordinate(self):
        rows = [
            ("He", "nsubj", 1, "PRON", "PRP"),
            ("said", "ROOT", 1, "VERB", "VBD"),
            ("that", "mark", 4, "SCONJ", "IN"),
            ("it", "nsubj", 4, "PRON", "PRP"),
            ("rained", "ccomp", 1, "VERB", "VBD"),
            ("and", "cc", 4, "CCONJ", "CC"),
            ("that", "mark", 8, "SCONJ", "IN"),
            ("it", "nsubj", 8, "PRON", "PRP"),
            ("snowed", "conj", 4, "VERB", "VBD"),
            (".", "punct", 1, "PUNCT", "."),
        ]
        sentence, view = tree(rows)
        tags = [
            ("He", "subject", 1),
            ("said", "verb", 1),
            ("and that it snowed", "coordinated_clause", 2),
        ]
        items, _ = generate(sentence, view, tags)
        for template in ("clause_boundary", "clause_type_classify"):
            item = by_template(items, template)[0]
            self.assertIn(R.CLAUSE_TYPE_MISMATCH, item["reasons"], template)

    def test_coordinated_predicate_is_not_coordinate_clause(self):
        rows = [
            ("They", "nsubj", 1, "PRON", "PRP"),
            ("sort", "ROOT", 1, "VERB", "VBP"),
            ("data", "dobj", 1, "NOUN", "NNS"),
            ("and", "cc", 1, "CCONJ", "CC"),
            ("usually", "advmod", 5, "ADV", "RB"),
            ("refer", "conj", 1, "VERB", "VBP"),
            ("to", "prep", 5, "ADP", "IN"),
            ("types", "pobj", 6, "NOUN", "NNS"),
            (".", "punct", 1, "PUNCT", "."),
        ]
        sentence, view = tree(rows)
        tags = [
            ("They sort data", "main_clause", 1),
            ("usually refer to types", "coordinated_clause", 1),
        ]
        items, _ = generate(sentence, view, tags)
        for template in ("clause_boundary", "clause_type_classify"):
            item = by_template(items, template)[0]
            self.assertIn(R.CLAUSE_TYPE_MISMATCH, item["reasons"], template)

    def test_for_clause_is_not_adverbial(self):
        rows = [
            ("We", "nsubj", 1, "PRON", "PRP"),
            ("stayed", "ROOT", 1, "VERB", "VBD"),
            ("for", "mark", 4, "SCONJ", "IN"),
            ("it", "nsubj", 4, "PRON", "PRP"),
            ("rained", "advcl", 1, "VERB", "VBD"),
            (".", "punct", 1, "PUNCT", "."),
        ]
        sentence, view = tree(rows)
        tags = [("We stayed", "main_clause", 1), ("for it rained", "cause_clause", 2)]
        items, _ = generate(sentence, view, tags)
        for template in ("clause_boundary", "clause_type_classify"):
            item = by_template(items, template)[0]
            self.assertIn(R.CLAUSE_TYPE_MISMATCH, item["reasons"], template)

    def test_s_quote_closes_a_pending_quoted_term(self):
        sentence = "They built 'megacities' and traded in 1990's markets."
        span = {"start": 11, "end": 22, "text": sentence[11:22]}
        closed = close_quotes(sentence, span, 0, len(sentence))
        self.assertEqual(closed["text"], "'megacities'")
        possessive = "Their students' books were new."
        span = {"start": 6, "end": 15, "text": possessive[6:15]}
        self.assertIs(close_quotes(possessive, span, 0, len(possessive)), span)

    def test_quote_starting_with_a_number_is_a_quote_mark(self):
        self.assertIsNone(single_quote_role("in the '60s and '70s", 7))
        sentence = "labels ranging from '15 volts (slight)' to '450 volts (danger)' were used."
        self.assertEqual(single_quote_role(sentence, sentence.index("'15")), "open")
        end = sentence.index("' were")
        span = {"start": 0, "end": end, "text": sentence[:end]}
        closed = close_quotes(sentence, span, 0, end + 1)
        self.assertEqual(
            closed["text"], "labels ranging from '15 volts (slight)' to '450 volts (danger)'"
        )

    def test_close_quotes_reattaches_closing_dash_within_raw_span(self):
        sentence = "Several processes - melting, refining - take place in the furnace."
        end = sentence.index(" - take")
        span = {"start": 0, "end": end, "text": sentence[:end]}
        closed = close_quotes(sentence, span, 0, sentence.index("take"))
        self.assertEqual(closed["text"], "Several processes - melting, refining -")
        self.assertIs(close_quotes(sentence, span, 0, end), span)  # The dash is outside the tag.
        blast = "The icy blast—the so-called katabatic wind—can reach 300 km/hr."
        end = blast.index("—can")
        span = {"start": 0, "end": end, "text": blast[:end]}
        self.assertEqual(close_quotes(blast, span, 0, end + 1)["text"], blast[: end + 1])
        # A range dash between digits is not one of the pair (review A5).
        ranged = "The 5–10 range – a useful convention – matters."
        end = ranged.index(" – a")
        span = {"start": 0, "end": end, "text": ranged[:end]}
        self.assertIs(close_quotes(ranged, span, 0, len(ranged)), span)
        # A digit on one side only does not make a range (review U6, 29_63_p3 A-03).
        share = "Some $265 billion of the increase - 93% of it - was ascribed to services."
        end = share.index(" - was")
        span = {"start": 0, "end": end, "text": share[:end]}
        closed = close_quotes(share, span, 0, len(share))
        self.assertEqual(closed["text"], "Some $265 billion of the increase - 93% of it -")
        year = "Sales in 2010–a turning point–doubled again."
        end = year.index("–doubled")
        span = {"start": 0, "end": end, "text": year[:end]}
        self.assertEqual(close_quotes(year, span, 0, len(year))["text"], year[: end + 1])

    def test_t1_distractor_outside_its_tag_span_is_downgraded(self):
        sentence, view = tree(
            [
                ("They", "nsubj", 1, "PRON", "PRP"),
                ("wanted", "ROOT", 1, "VERB", "VBD"),
                ("to", "aux", 3, "PART", "TO"),
                ("close", "xcomp", 1, "VERB", "VB"),
                ("the", "det", 5, "DET", "DT"),
                ("gate", "dobj", 3, "NOUN", "NN"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        tags = [
            ("They", "subject", 1),
            ("wanted", "predicate", 1),
            ("close the gate", "infinitive_complement", 2),
        ]
        items, _ = generate(sentence, view, tags)
        t1 = by_template(items, "main_verb_tap")[0]
        self.assertEqual([s["text"] for s in t1["answer"]["spans"]], ["wanted"])
        self.assertEqual([d["text"] for d in t1["distractors"]], ["to close"])
        self.assertIn(R.VERB_UNIT_OUTSIDE_SPAN, t1["reasons"])
        self.assertEqual(t1["status"], "needs_review")
        finite_items = by_template(items, "finite_or_not")
        main = next(i for i in finite_items if i["evidence"]["variant"] == "main")
        self.assertEqual(main["status"], "auto_pass")  # The root's own unit is inside its tag.
        tags[2] = ("to close the gate", "infinitive_complement", 2)
        t1 = by_template(generate(sentence, view, tags)[0], "main_verb_tap")[0]
        self.assertEqual(t1["status"], "auto_pass")


class PredicativeClauseTests(unittest.TestCase):
    """Review V1: a because clause that is be's complement is not an adverbial clause."""

    def reason_tree(self, middle):
        rows = [("That", "nsubj", 2, "PRON", "DT"), ("could", "aux", 2, "AUX", "MD")]
        rows += [("be", "ROOT", 2, "AUX", "VB")] + middle
        start = len(rows)
        rows += [
            ("because", "mark", start + 2, "SCONJ", "IN"),
            ("ice", "nsubj", start + 2, "NOUN", "NN"),
            ("melted", "advcl", 2, "VERB", "VBD"),
            (".", "punct", 2, "PUNCT", "."),
        ]
        return tree(rows)

    def clause_items(self, sentence, view, predicate):
        tags = [
            ("That", "subject", 1),
            (predicate, "verb_phrase", 1),
            ("because ice melted", "reason_clause", 2),
        ]
        items, _ = generate(sentence, view, tags)
        return [
            i
            for i in items
            if i["template"] in {"clause_boundary", "clause_type_classify", "adverbial_relation"}
        ]

    def test_because_clause_after_bare_be_is_not_auto_pass(self):
        # "The researchers say that could be because many of those artefacts have ..." (H-04);
        # an adverb such as "only" may stand between ("it's only because ...").
        sentence, view = self.reason_tree([("only", "advmod", 2, "ADV", "RB")])
        items = self.clause_items(sentence, view, "could be")
        self.assertEqual(len(items), 3)
        for item in items:
            self.assertEqual(item["status"], "needs_review", item["template"])
        sentence, view = self.reason_tree([])
        items = self.clause_items(sentence, view, "could be")
        self.assertEqual(
            {i["template"] for i in items},
            {"clause_boundary", "clause_type_classify", "adverbial_relation"},
        )
        for item in items:
            self.assertEqual(item["status"], "needs_review", item["template"])
            self.assertIn(R.CLAUSE_TYPE_MISMATCH, item["reasons"])

    def test_because_clause_after_be_with_complement_stays_adverbial(self):
        sentence, view = self.reason_tree([("cold", "acomp", 2, "ADJ", "JJ")])
        items = self.clause_items(sentence, view, "could be cold")
        self.assertEqual(len(items), 3)
        for item in items:
            self.assertNotIn(R.CLAUSE_TYPE_MISMATCH, item["reasons"], item["template"])
            self.assertEqual(item["status"], "auto_pass", item["template"])

    def test_because_clause_after_other_verb_stays_adverbial(self):
        sentence, view = tree(
            [
                ("That", "nsubj", 1, "PRON", "DT"),
                ("failed", "ROOT", 1, "VERB", "VBD"),
                ("because", "mark", 4, "SCONJ", "IN"),
                ("ice", "nsubj", 4, "NOUN", "NN"),
                ("melted", "advcl", 1, "VERB", "VBD"),
                (".", "punct", 1, "PUNCT", "."),
            ]
        )
        items = self.clause_items(sentence, view, "failed")
        self.assertEqual(len(items), 3)
        for item in items:
            self.assertEqual(item["status"], "auto_pass", item["template"])

    def test_fronted_clause_on_bare_be_stays_adverbial(self):
        # "When the opponent's king is checkmated, the game is over, ..." (over parsed as advmod).
        sentence, view = tree(
            [
                ("When", "mark", 2, "SCONJ", "WRB"),
                ("ice", "nsubj", 2, "NOUN", "NN"),
                ("melted", "advcl", 5, "VERB", "VBD"),
                (",", "punct", 5, "PUNCT", ","),
                ("it", "nsubj", 5, "PRON", "PRP"),
                ("was", "ROOT", 5, "AUX", "VBD"),
                ("over", "advmod", 5, "ADV", "RB"),
                (".", "punct", 5, "PUNCT", "."),
            ]
        )
        tags = [
            ("When ice melted", "time_clause", 2),
            ("it", "subject", 1),
            ("was over", "verb_phrase", 1),
        ]
        items, _ = generate(sentence, view, tags)
        clause = [i for i in items if i["template"] in {"clause_boundary", "adverbial_relation"}]
        self.assertEqual(len(clause), 2)
        for item in clause:
            self.assertNotIn(R.CLAUSE_TYPE_MISMATCH, item["reasons"], item["template"])


MAIN_ROWS = [
    ("The", "det", 1, "DET", "DT"),
    ("plans", "nsubj", 4, "NOUN", "NNS"),
    ("may", "aux", 4, "AUX", "MD"),
    ("not", "neg", 4, "PART", "RB"),
    ("work", "ROOT", 4, "VERB", "VB"),
    ("out", "prt", 4, "ADP", "RP"),
    ("when", "mark", 8, "SCONJ", "IN"),
    ("prices", "nsubj", 8, "NOUN", "NNS"),
    ("rise", "advcl", 4, "VERB", "VBP"),
    ("today", "npadvmod", 8, "NOUN", "NN"),
    (".", "punct", 4, "PUNCT", "."),
]


class SentenceReviewTests(unittest.TestCase):
    """Round 7: defective source sentences and list items never pass automatically."""

    def test_review_reasons(self):
        points = ["原句问题：主语是复数，规范写法是 rise。", "when 引导时间状语从句。"]
        self.assertEqual(review_reasons("The plans may not work out.", points), [R.SOURCE_QUALITY])
        # Only a note opening with the prefix counts.
        self.assertEqual(review_reasons("The plans may not work out.", ["见原句问题：无"]), [])
        for item in (
            "- is calm, and is not excitable when something unexpected happens",
            "• The S.K.I. Bicycle Courier Service first started in the Sudan.",
            "* peat: a brown deposit formed by the partial decomposition of vegetation",
            "1. Valid assessment considers cultural and linguistic diversity.",
            "2) Restriction from participation shall continue until an evaluation.",
            "100. The last rule applies to every member.",
            "a) The first option is cheaper.",
            "(iv) The fourth option is slower.",
            "· The plans may not work out.",
            "− The plans may not work out.",
        ):
            self.assertEqual(review_reasons(item), [R.LIST_ITEM], item)
        for sentence in (
            "1.5 million people moved to the city in 2010.",
            "In 5–10 years the plans may work out.",
            "1990s cities grew fast.",
            "-based materials were used.",
            "— We agreed, she said.",  # An em dash can open dialogue.
            "e.g. the plans may not work out.",
        ):
            self.assertEqual(review_reasons(sentence), [], sentence)
        self.assertEqual(
            review_reasons("- The plans may not work out.", points[:1]),
            [R.SOURCE_QUALITY, R.LIST_ITEM],
        )

    def test_letter_labels_need_a_neighbouring_letter(self):
        passage = [
            "A. Participation by individuals with Down Syndrome is limited.",
            "B. SOC requires temporary restriction of individuals from participation.",
            "A. K. Lakhina says the government recognises the potential.",
            "K. Anders Ericsson thinks anyone can become an expert.",
        ]
        letters = list_letters(passage)
        self.assertEqual(letters, frozenset({"A", "B"}))
        self.assertEqual(review_reasons(passage[1], (), letters), [R.LIST_ITEM])
        self.assertEqual(review_reasons(passage[0], (), letters), [R.LIST_ITEM])
        self.assertEqual(review_reasons(passage[2], (), letters), [])  # Initials, not a label.
        self.assertEqual(review_reasons(passage[3], (), letters), [])
        self.assertEqual(list_letters(passage[2:]), frozenset())
        self.assertEqual(review_reasons(passage[1]), [])  # Needs the passage's letters.

    def test_review_reasons_send_every_item_to_review(self):
        sentence, view = tree(MAIN_ROWS)
        spans = locate_spans(
            sentence, [{"text": t, "role": r, "level": lv} for t, r, lv in main_tags()]
        )
        plain, _ = generate_templates(source(sentence), spans, view)
        flagged, _ = generate_templates(source(sentence), spans, view, [R.SOURCE_QUALITY])
        self.assertIn("auto_pass", {i["status"] for i in plain})
        self.assertEqual([i["item_id"] for i in flagged], [i["item_id"] for i in plain])
        for item in flagged:
            self.assertEqual(item["status"], "needs_review")
            self.assertIn(R.SOURCE_QUALITY, item["reasons"])

    def test_report_counts_sentence_review_items(self):
        first, second = source("x"), source("y", sid="A-02")
        items = [
            {"template": "main_verb_tap", "source": first, "reasons": [R.SOURCE_QUALITY]},
            {
                "template": "finite_or_not",
                "source": first,
                "reasons": [R.AMBIGUOUS_SPAN_LOCATION, R.SOURCE_QUALITY],
            },
            {
                "template": "main_verb_tap",
                "source": second,
                "reasons": [R.LIST_ITEM, R.SOURCE_QUALITY],
            },
        ]
        stats = sentence_review_stats(items)
        self.assertEqual(
            stats[R.SOURCE_QUALITY],
            {
                "sentences": 2,
                "items_by_template": {"finite_or_not": 1, "main_verb_tap": 2},
                "only_reason": 2,
            },
        )
        self.assertEqual(
            stats[R.LIST_ITEM],
            {"sentences": 1, "items_by_template": {"main_verb_tap": 1}, "only_reason": 1},
        )

    def test_collect_reads_notes_and_list_markers(self):
        sentence, view = tree(MAIN_ROWS)
        listed, listed_view = tree(
            [("-", "punct", 5, "PUNCT", ":")] + [(t, d, h + 1, p, g) for t, d, h, p, g in MAIN_ROWS]
        )
        structure = [{"text": t, "role": r, "level": lv} for t, r, lv in main_tags()]
        note = "when 引导时间状语从句。"
        rows = [
            ("A-01", sentence, [note, "may not 是情态动词加否定。"]),
            ("A-02", sentence, [note, "原句问题：句末缺少下文。"]),
            ("A-03", listed, [note, "may not 是情态动词加否定。"]),
        ]
        passage = {
            "source_kind": "reading_test",
            "test_id": "test",
            "passage_id": "test_p1",
            "sentences": [
                {"id": sid, "sentence": text, "structure": structure, "difficult_points": points}
                for sid, text, points in rows
            ],
        }

        class Cache:
            def parse(self, text, nlp):
                return {sentence: view, listed: listed_view}[text]

        with tempfile.TemporaryDirectory() as folder:
            data = Path(folder) / "data/reading_study"
            data.mkdir(parents=True)
            (data / "test_p1.json").write_text(json.dumps(passage), encoding="utf-8")
            candidates, rejections, _ = collect(Path(folder), None, Cache())
        self.assertFalse([r for r in rejections if r["scope"] == "sentence"])
        by_sentence = {}
        for item in candidates:
            by_sentence.setdefault(item["source"]["sentence_id"], []).append(item)
        self.assertIn("auto_pass", {i["status"] for i in by_sentence["A-01"]})
        for sid, reason in (("A-02", R.SOURCE_QUALITY), ("A-03", R.LIST_ITEM)):
            self.assertEqual(len(by_sentence[sid]), len(by_sentence["A-01"]))
            for item in by_sentence[sid]:
                self.assertEqual(item["status"], "needs_review")
                self.assertIn(reason, item["reasons"])


if __name__ == "__main__":
    unittest.main()
